from __future__ import annotations

import io
import json
import logging
import mimetypes
import os
import re
from datetime import datetime, timezone
from pathlib import Path
from typing import List, Optional
from urllib.parse import urlparse

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from pydantic import BaseModel, Field
from sqlalchemy import and_, text
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool

import models
from deps import (
    call_ai,
    call_ai_async,
    get_current_user,
    get_db,
)
from services.document_processor import extract_text_from_pdf_detailed
from services.api_key_pool import ApiKeyPoolExhausted
from services.storage_service import StorageService
from tutor.contract import tutor_contract_instruction
from uid_utils import resolve_by_id_or_uid

CHAT_UPLOAD_DIR = Path("uploads/chat_images")
CHAT_UPLOAD_DIR.mkdir(parents=True, exist_ok=True)

_SUPPORTED_IMAGE_TYPES = {"image/jpeg", "image/png", "image/gif", "image/webp"}
_SUPPORTED_DOCX_TYPES = {"application/vnd.openxmlformats-officedocument.wordprocessingml.document"}
_SUPPORTED_TEXT_EXTENSIONS = (".txt", ".md", ".markdown")
_SUPPORTED_TEXT_TYPES = {"text/plain", "text/markdown", "text/x-markdown"}
_MAX_IMAGE_BYTES = 20 * 1024 * 1024
_MAX_IMAGES_PER_MESSAGE = 10
_TUTOR_VISIBLE_OPTION_RE = re.compile(r"(?im)^\s*([A-F])[\).:-]\s+(.{2,220})\s*$")
_TUTOR_ANSWER_FIELD_RE = re.compile(r'(?is)"answer"\s*:\s*"(.*?)"\s*,\s*"tutor_state"\s*:')
_TUTOR_STATE_FIELD_RE = re.compile(r'(?is)"tutor_state"\s*:\s*(\{.*?\})\s*,\s*"options"\s*:')
_TUTOR_OPTIONS_FIELD_RE = re.compile(r'(?is)"options"\s*:\s*(\[.*?\])')
_INTERNAL_GRAPH_GUIDANCE_MARKERS = [
    "if a visual would materially improve understanding,",
    "prefer ```graphjson for this response",
    "for `graphjson`, use schema:",
    "do not include any graph or diagram block unless the user explicitly asks for one.",
]
_PERSON_IDENTITY_REQUEST_RE = re.compile(
    r"\b(?:"
    r"who\s+(?:is|are)\s+(?:this|that|he|she|they|the\s+(?:person|man|woman|boy|girl))"
    r"|identify\s+(?:this|that|the)\s+(?:person|man|woman|boy|girl)"
    r"|(?:what(?:'s| is)|tell\s+me)\s+(?:his|her|their|the\s+person'?s)\s+name"
    r")\b",
    re.IGNORECASE,
)

logger = logging.getLogger(__name__)
router = APIRouter(prefix="/api", tags=["chat"])


def _usage_limit_payload(exc: ApiKeyPoolExhausted) -> dict:
    seconds = exc.reset_after_seconds
    return {
        "detail": "You've reached your usage limit.",
        "code": "ai_provider_limit_exceeded",
        "provider": exc.provider or "ai",
        "reset_at": exc.reset_at,
        "reset_after_seconds": seconds,
        "reset_after_hours": round(seconds / 3600, 1) if seconds is not None else None,
    }


def _raise_if_usage_limit_error(exc: Exception) -> None:
    current = exc
    seen = set()
    while current and id(current) not in seen:
        seen.add(id(current))
        if isinstance(current, ApiKeyPoolExhausted):
            payload = _usage_limit_payload(current)
            headers = {"X-AI-Limit-Code": "ai_provider_limit_exceeded"}
            if current.reset_at:
                headers["X-AI-Limit-Reset"] = str(current.reset_at)
            if current.reset_after_seconds is not None:
                headers["Retry-After"] = str(max(1, int(current.reset_after_seconds)))
                headers["X-AI-Limit-Reset-After"] = str(max(0, int(current.reset_after_seconds)))
            raise HTTPException(status_code=429, detail=payload, headers=headers)
        current = current.__cause__ or current.__context__


def _is_usage_limit_error(exc: Exception) -> bool:
    current = exc
    seen = set()
    while current and id(current) not in seen:
        seen.add(id(current))
        if isinstance(current, ApiKeyPoolExhausted):
            return True
        current = current.__cause__ or current.__context__
    return False


def _load_test_fallback_enabled(user: models.User) -> bool:
    identifiers = {
        item.strip().lower()
        for item in os.getenv("AI_LOAD_TEST_FALLBACK_USERS", "").split(",")
        if item.strip()
    }
    return any(
        value and str(value).strip().lower() in identifiers
        for value in (getattr(user, "id", None), getattr(user, "username", None), getattr(user, "email", None))
    )


def _load_test_answer(question: str) -> str:
    prompt = (question or "this topic").strip()
    return (
        "Active recall helps memory because it makes you retrieve an idea instead of only rereading it. "
        f"For this load-test prompt, focus on the core idea, a quick example, and one self-check: {prompt[:180]}"
    )


def _normalized_upload_mime(content_type: Optional[str], filename: str) -> str:
    mime = (content_type or "").split(";", 1)[0].strip().lower()
    if mime == "image/jpg":
        return "image/jpeg"

    guessed, _ = mimetypes.guess_type(filename or "")
    if not mime or mime == "application/octet-stream":
        mime = (guessed or "").lower()
    return "image/jpeg" if mime == "image/jpg" else mime


def _is_person_identity_request(question: str) -> bool:
    return bool(_PERSON_IDENTITY_REQUEST_RE.search(question or ""))


class TutorOptionPayload(BaseModel):
    label: Optional[str] = None
    text: str

class TutorStatePayload(BaseModel):
    level: str = "intermediate"
    phase: str = "teach"
    verdict: str = "not_applicable"
    confidence: float = 0.65
    objective: str = "Build understanding step by step"
    next_action: str = "Try the next small step"
    hint_level: int = 2
    current_step: int = 1
    total_steps: int = 0
    expected_step_answer: Optional[str] = None
    final_answer: Optional[str] = None
    skills_used: List[str] = Field(default_factory=list)
    misconceptions: List[str] = Field(default_factory=list)
    mastery_score: float = 0.0
    correct_streak: int = 0
    wrong_streak: int = 0
    lesson_plan: Optional[dict] = None

class TutorResponsePayload(BaseModel):
    answer: str
    tutor_state: Optional[TutorStatePayload] = None
    options: List[TutorOptionPayload] = Field(default_factory=list)

def _normalize_tutor_options(raw_options: object) -> list[dict]:
    if not isinstance(raw_options, list):
        return []

    normalized = []
    for index, option in enumerate(raw_options[:6]):
        fallback_label = chr(ord("A") + index)
        if isinstance(option, dict):
            text = str(option.get("text") or option.get("label") or option.get("value") or "").strip()
            label = str(option.get("label") or fallback_label).strip()
        else:
            text = str(option or "").strip()
            match = re.match(r"^\s*([A-F])[\).:-]\s*(.+)$", text, flags=re.I)
            label = match.group(1).upper() if match else fallback_label
            text = match.group(2).strip() if match else text

        if not text:
            continue
        label = label[:1].upper() if label else fallback_label
        normalized.append({
            "id": label,
            "label": label,
            "text": text,
            "value": f"{label}. {text}",
        })
    return normalized

def _extract_visible_tutor_options(response_text: str) -> list[dict]:
    raw = response_text or ""
    visible_options = [
        f"{match.group(1).upper()}. {match.group(2).strip()}"
        for match in _TUTOR_VISIBLE_OPTION_RE.finditer(raw)
    ]
    has_mcq_language = re.search(
        r"\b(mcq|multiple choice|choose|which option|select one|quick check)\b",
        raw,
        flags=re.I,
    )
    if len(visible_options) >= 3 or (len(visible_options) >= 2 and has_mcq_language):
        return _normalize_tutor_options(visible_options)
    return []

def _strip_internal_graph_guidance(text: str) -> str:
    raw = str(text or "").replace("\r\n", "\n")
    if not raw:
        return ""
    lowered = raw.lower()
    cut_at = -1
    for marker in _INTERNAL_GRAPH_GUIDANCE_MARKERS:
        index = lowered.find(marker)
        if index >= 0 and (cut_at == -1 or index < cut_at):
            cut_at = index
    if cut_at == -1:
        return raw.strip()
    return raw[:cut_at].strip()

def _fallback_chat_title(text: str) -> str:
    raw = _strip_internal_graph_guidance(text or "")
    raw = re.sub(r"```.*?```", " ", raw, flags=re.S)
    raw = re.sub(r"https?://\S+", " ", raw)
    raw = re.sub(r"[*_`#>\[\]{}()|]", " ", raw)
    raw = re.sub(r"\s+", " ", raw).strip(" .,:;!?")
    words = raw.split()
    if not words:
        return "New Chat"
    title = " ".join(words[:6])
    if len(words) > 6:
        title += "..."
    return title[:50]

def _apply_fallback_chat_title(chat_session, prompt_text: str) -> None:
    if chat_session and chat_session.title == "New Chat":
        title = _fallback_chat_title(prompt_text)
        if title != "New Chat":
            chat_session.title = title

def _pydantic_to_dict(model: BaseModel | None) -> dict:
    if not model:
        return {}
    if hasattr(model, "model_dump"):
        return model.model_dump()
    return model.dict()

def _strip_legacy_tutor_markers(text: str) -> str:
    return re.sub(r"(?im)^\s*TUTOR_(?:OPTIONS|STATE)\s*:.*$", "", text or "").strip()

def _strip_json_code_fence(text: str) -> str:
    raw = (text or "").strip()
    if raw.startswith("```"):
        raw = raw.split("\n", 1)[-1]
        if raw.rstrip().endswith("```"):
            raw = raw.rsplit("```", 1)[0]
    return raw.strip()

def _normalize_tutor_jsonish_text(text: str) -> str:
    return (
        str(text or "")
        .replace("\u201c", '"')
        .replace("\u201d", '"')
        .replace("\u2018", "'")
        .replace("\u2019", "'")
    )

def _escape_latex_backslashes_for_json(text: str) -> str:
    return re.sub(r"(?<!\\)\\(?=[A-Za-z])", r"\\\\", text or "")

def _json_decode_lenient(value: str):
    raw = _normalize_tutor_jsonish_text(_strip_json_code_fence(value))
    escaped = _escape_latex_backslashes_for_json(raw)
    candidates = [escaped]
    if escaped != raw:
        candidates.append(raw)

    for candidate in candidates:
        try:
            return json.loads(candidate)
        except Exception:
            continue
    return None

def _decode_jsonish_string(value: str) -> str:
    escaped = _escape_latex_backslashes_for_json(value or "")
    try:
        decoded = json.loads(f'"{escaped}"')
        return str(decoded).strip()
    except Exception:
        return (value or "").replace('\\"', '"').replace("\\\\", "\\").strip()

def _coerce_tutor_answer(value: object) -> str:
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, list):
        return "\n".join(
            str(item or "").strip()
            for item in value
            if str(item or "").strip()
        ).strip()
    return ""

def _coerce_pseudo_tutor_answer_array(value: str) -> str:
    inner = re.sub(r"^\s*\[|\]\s*$", "", value or "").strip()
    lines = []
    for raw_line in inner.splitlines():
        line = raw_line.strip().rstrip(",").strip()
        if not line:
            continue
        had_bullet = bool(re.match(r"^[*\-]\s+", line))
        line = re.sub(r"^[*\-]\s+", "", line).strip().strip('"').strip("'").strip()
        if not line:
            continue
        lines.append(f"- {line}" if had_bullet else line)
    return "\n".join(lines).strip()

def _trim_tutor_text(value: object, fallback: str, max_len: int = 120) -> str:
    text = re.sub(r"\s+", " ", str(value or "").strip())
    return (text or fallback)[:max_len]

def _bounded_int(value: object, default: int = 0, min_value: int = 0, max_value: int = 999) -> int:
    try:
        parsed = int(value)
    except (TypeError, ValueError):
        parsed = default
    return max(min_value, min(max_value, parsed))

def _bounded_float(value: object, default: float = 0.0) -> float:
    try:
        parsed = float(value)
    except (TypeError, ValueError):
        parsed = default
    return max(0.0, min(1.0, parsed))

def _trim_string_list(value: object, max_items: int = 8, max_len: int = 80) -> list[str]:
    if not isinstance(value, list):
        return []
    items = []
    for item in value[:max_items]:
        text = re.sub(r"\s+", " ", str(item or "").strip())
        if text:
            items.append(text[:max_len])
    return items

def _normalize_weakness_topic(value: object, fallback: str = "") -> str:
    text_value = re.sub(r"\s+", " ", str(value or "").strip())
    if not text_value:
        text_value = fallback
    text_value = re.sub(r"^(misconception|mistake|error|gap)\s*[:\-]\s*", "", text_value, flags=re.I)
    return text_value[:255]

def _record_tutor_weakness_signals(
    db: Session,
    user_id: int,
    tutor_state: Optional[dict],
    tutor_choice: Optional[str] = None,
) -> list[str]:
    if not tutor_state:
        return []

    verdict = str(tutor_state.get("verdict") or "").strip().lower()
    if verdict not in {"partly_correct", "not_yet"}:
        return []

    objective = _normalize_weakness_topic(tutor_state.get("objective"), "AI tutor mistake")
    skills = _trim_string_list(tutor_state.get("skills_used"), 4, 80)
    misconceptions = _trim_string_list(tutor_state.get("misconceptions"), 5, 120)
    raw_topics = misconceptions or skills or [objective]
    topics: list[str] = []
    for raw_topic in raw_topics:
        topic = _normalize_weakness_topic(raw_topic, objective)
        if topic and topic.lower() not in {item.lower() for item in topics}:
            topics.append(topic)

    now = datetime.now(timezone.utc)
    recorded: list[str] = []
    for topic in topics[:5]:
        existing = db.query(models.UserWeakArea).filter(
            models.UserWeakArea.user_id == user_id,
            models.UserWeakArea.topic == topic,
        ).first()
        penalty = 18 if verdict == "not_yet" else 10
        if existing:
            existing.subtopic = existing.subtopic or objective
            existing.total_questions = (existing.total_questions or 0) + 1
            existing.incorrect_count = (existing.incorrect_count or 0) + 1
            existing.consecutive_wrong = (existing.consecutive_wrong or 0) + 1
            existing.last_wrong_streak = max(existing.last_wrong_streak or 0, existing.consecutive_wrong or 0)
            existing.accuracy = round(((existing.correct_count or 0) / max(existing.total_questions or 0, 1)) * 100, 1)
            existing.weakness_score = min(100.0, max(existing.weakness_score or 0.0, 35.0) + penalty)
            existing.priority = max(existing.priority or 0, 8 if verdict == "not_yet" else 6)
            existing.status = "needs_practice"
            existing.last_updated = now
        else:
            db.add(models.UserWeakArea(
                user_id=user_id,
                topic=topic,
                subtopic=objective,
                total_questions=1,
                correct_count=0,
                incorrect_count=1,
                accuracy=0.0,
                weakness_score=55.0 if verdict == "not_yet" else 42.0,
                consecutive_wrong=1,
                last_wrong_streak=1,
                status="needs_practice",
                priority=8 if verdict == "not_yet" else 6,
                first_identified=now,
                last_updated=now,
            ))
        recorded.append(topic)

    if recorded:
        logger.info(
            "Recorded tutor weakness signal: user_id=%s verdict=%s topics=%s choice=%s",
            user_id,
            verdict,
            recorded,
            (tutor_choice or "")[:80],
        )
    return recorded

def _normalize_lesson_plan(value: object) -> Optional[dict]:
    if not isinstance(value, dict):
        return None
    steps = value.get("steps") if isinstance(value.get("steps"), list) else []
    normalized_steps = []
    for index, step in enumerate(steps[:12], start=1):
        if not isinstance(step, dict):
            continue
        normalized_steps.append({
            "id": _bounded_int(step.get("id"), index, 1, 99),
            "title": _trim_tutor_text(step.get("title"), f"Step {index}", 80),
            "expected": _trim_tutor_text(step.get("expected"), "", 240),
            "skill": _trim_tutor_text(step.get("skill"), "", 80),
            "misconception": _trim_tutor_text(step.get("misconception"), "", 120),
        })
    return {
        "goal": _trim_tutor_text(value.get("goal"), "Build understanding step by step", 140),
        "steps": normalized_steps,
        "current_step": _bounded_int(value.get("current_step"), 1, 1, max(1, len(normalized_steps) or 99)),
        "final_answer": _trim_tutor_text(value.get("final_answer"), "", 500),
    }

def _normalize_tutor_state(
    raw_state: object,
    reply_style: str = "guided",
    tutor_choice: Optional[str] = None,
) -> dict:
    data = raw_state if isinstance(raw_state, dict) else {}
    style = (reply_style or "guided").strip().lower()

    levels = {"beginner", "intermediate", "advanced"}
    phases = {"diagnose", "teach", "practice", "check", "review"}
    verdicts = {"correct", "partly_correct", "not_yet", "needs_attempt", "not_applicable"}
    default_phase = {
        "hint": "teach",
        "guided": "teach",
        "check": "check",
        "quiz": "practice",
    }.get(style, "teach")

    level = str(data.get("level") or "intermediate").strip().lower()
    phase = str(data.get("phase") or default_phase).strip().lower()
    verdict = str(data.get("verdict") or ("needs_attempt" if tutor_choice else "not_applicable")).strip().lower()
    level = {"introductory": "beginner"}.get(level, level)
    phase = {"exploration": "teach"}.get(phase, phase)
    verdict = {"in_progress": "not_applicable", "pending": "needs_attempt"}.get(verdict, verdict)

    try:
        confidence = float(data.get("confidence", 0.65))
    except (TypeError, ValueError):
        confidence = 0.65
    confidence = max(0.0, min(1.0, confidence))

    try:
        hint_level = int(data.get("hint_level", 1 if style == "hint" else 2))
    except (TypeError, ValueError):
        hint_level = 2
    hint_level = max(1, min(3, hint_level))

    lesson_plan = _normalize_lesson_plan(data.get("lesson_plan"))
    total_steps = _bounded_int(
        data.get("total_steps") or (len(lesson_plan.get("steps", [])) if lesson_plan else 0),
        0,
        0,
        99,
    )
    current_step = _bounded_int(
        data.get("current_step") or (lesson_plan.get("current_step") if lesson_plan else 1),
        1,
        1,
        max(1, total_steps or 99),
    )

    return {
        "level": level if level in levels else "intermediate",
        "phase": phase if phase in phases else default_phase,
        "verdict": verdict if verdict in verdicts else "not_applicable",
        "confidence": round(confidence, 2),
        "objective": _trim_tutor_text(data.get("objective"), "Build understanding step by step", 96),
        "next_action": _trim_tutor_text(data.get("next_action"), "Try the next small step", 120),
        "hint_level": hint_level,
        "current_step": current_step,
        "total_steps": total_steps,
        "expected_step_answer": _trim_tutor_text(data.get("expected_step_answer"), "", 300),
        "final_answer": _trim_tutor_text(data.get("final_answer"), "", 500),
        "skills_used": _trim_string_list(data.get("skills_used"), 8, 80),
        "misconceptions": _trim_string_list(data.get("misconceptions"), 10, 120),
        "mastery_score": round(_bounded_float(data.get("mastery_score"), 0.0), 2),
        "correct_streak": _bounded_int(data.get("correct_streak"), 0, 0, 99),
        "wrong_streak": _bounded_int(data.get("wrong_streak"), 0, 0, 99),
        "lesson_plan": lesson_plan,
    }

def _extract_json_object(text: str) -> Optional[dict]:
    raw = _normalize_tutor_jsonish_text(_strip_json_code_fence(text))
    escaped = _escape_latex_backslashes_for_json(raw)
    candidates = [escaped]
    if escaped != raw:
        candidates.append(raw)

    for candidate in candidates:
        decoder = json.JSONDecoder()
        for index, char in enumerate(candidate):
            if char != "{":
                continue
            try:
                obj, _ = decoder.raw_decode(candidate[index:])
            except Exception:
                continue
            if isinstance(obj, dict):
                if "answer" in obj or "tutor_state" in obj:
                    return obj
    return None

def _extract_tutor_contract_fallback(
    response_text: str,
    reply_style: str = "guided",
    tutor_choice: Optional[str] = None,
) -> Optional[tuple[str, list[dict], dict]]:
    raw = _normalize_tutor_jsonish_text(_strip_json_code_fence(response_text))
    if '"answer"' not in raw or '"tutor_state"' not in raw:
        return None

    answer_array_match = re.search(r'(?is)"answer"\s*:\s*(\[[\s\S]*?\])\s*,\s*"tutor_state"\s*:', raw)
    answer_match = _TUTOR_ANSWER_FIELD_RE.search(raw)
    answer_value = ""
    if answer_array_match:
        answer_value = (
            _coerce_tutor_answer(_json_decode_lenient(answer_array_match.group(1)))
            or _coerce_pseudo_tutor_answer_array(answer_array_match.group(1))
        )
    elif answer_match:
        answer_value = _decode_jsonish_string(answer_match.group(1))
    if not answer_value:
        return None

    tutor_state = {}
    state_match = _TUTOR_STATE_FIELD_RE.search(raw)
    if state_match:
        decoded_state = _json_decode_lenient(state_match.group(1))
        if isinstance(decoded_state, dict):
            tutor_state = decoded_state

    tutor_options = []
    options_match = _TUTOR_OPTIONS_FIELD_RE.search(raw)
    if options_match:
        decoded_options = _json_decode_lenient(options_match.group(1))
        tutor_options = _normalize_tutor_options(decoded_options)

    return (
        answer_value,
        tutor_options,
        _normalize_tutor_state(tutor_state, reply_style=reply_style, tutor_choice=tutor_choice),
    )

def _parse_tutor_response(
    response_text: str,
    reply_style: str = "guided",
    tutor_choice: Optional[str] = None,
) -> tuple[str, list[dict], dict]:
    raw = response_text or ""
    parsed = _extract_json_object(raw)
    if parsed:
        try:
            payload = TutorResponsePayload(**parsed)
            answer = (payload.answer or "").strip()
            tutor_state = _normalize_tutor_state(
                _pydantic_to_dict(payload.tutor_state),
                reply_style=reply_style,
                tutor_choice=tutor_choice,
            )
            tutor_options = _normalize_tutor_options([
                _pydantic_to_dict(option)
                for option in (payload.options or [])
            ])
            return answer, tutor_options, tutor_state
        except Exception as exc:
            logger.warning("TutorResponse JSON validation failed: %s", exc)
            answer = _coerce_tutor_answer(parsed.get("answer")) if isinstance(parsed, dict) else ""
            tutor_state = _normalize_tutor_state(
                parsed.get("tutor_state") if isinstance(parsed, dict) else {},
                reply_style=reply_style,
                tutor_choice=tutor_choice,
            )
            tutor_options = _normalize_tutor_options(
                parsed.get("options") if isinstance(parsed, dict) else []
            )
            if answer:
                return answer, tutor_options, tutor_state

    contract_fallback = _extract_tutor_contract_fallback(raw, reply_style, tutor_choice)
    if contract_fallback:
        return contract_fallback

    clean_answer = _strip_legacy_tutor_markers(_strip_json_code_fence(raw))
    return (
        clean_answer,
        _extract_visible_tutor_options(clean_answer),
        _normalize_tutor_state({}, reply_style=reply_style, tutor_choice=tutor_choice),
    )

def _attempt_evaluation_to_dict(value: object) -> dict:
    if not value:
        return {}
    if isinstance(value, dict):
        return value
    return {
        "verdict": getattr(value, "verdict", None),
        "confidence": getattr(value, "confidence", None),
        "rationale": getattr(value, "rationale", None),
        "expected_answer": getattr(value, "expected_answer", None),
        "next_action": getattr(value, "next_action", None),
        "is_final_answer": getattr(value, "is_final_answer", None),
        "final_answer_correct": getattr(value, "final_answer_correct", None),
        "misconception": getattr(value, "misconception", None),
    }

def _tutor_plan_to_dict(value: object) -> dict:
    if not value:
        return {}
    if isinstance(value, dict):
        return value
    return {
        "goal": getattr(value, "goal", None),
        "current_step": getattr(value, "current_step", None),
        "total_steps": getattr(value, "total_steps", None),
        "steps": getattr(value, "steps", None),
        "expected_step_answer": getattr(value, "expected_step_answer", None),
        "final_answer": getattr(value, "final_answer", None),
        "skills_used": getattr(value, "skills_used", None),
        "misconceptions": getattr(value, "misconceptions", None),
        "mastery_score": getattr(value, "mastery_score", None),
    }

def _apply_tutor_plan(tutor_state: Optional[dict], tutor_plan: object) -> Optional[dict]:
    plan_data = _normalize_lesson_plan(_tutor_plan_to_dict(tutor_plan))
    if not plan_data:
        return tutor_state
    state = dict(tutor_state or {})
    state["lesson_plan"] = plan_data
    state["objective"] = state.get("objective") or plan_data.get("goal")
    state["current_step"] = state.get("current_step") or plan_data.get("current_step") or 1
    state["total_steps"] = state.get("total_steps") or len(plan_data.get("steps", []))
    state["final_answer"] = state.get("final_answer") or plan_data.get("final_answer") or ""
    plan_source = _tutor_plan_to_dict(tutor_plan)
    state["skills_used"] = state.get("skills_used") or _trim_string_list(plan_source.get("skills_used"), 10, 80)
    state["misconceptions"] = state.get("misconceptions") or _trim_string_list(plan_source.get("misconceptions"), 12, 120)
    state["mastery_score"] = state.get("mastery_score", plan_source.get("mastery_score") or 0.0)
    current_step = _bounded_int(state.get("current_step"), 1, 1, max(1, state.get("total_steps") or 99))
    step = next((item for item in plan_data.get("steps", []) if item.get("id") == current_step), None)
    if step and not state.get("expected_step_answer"):
        state["expected_step_answer"] = step.get("expected", "")
    return state

def _apply_attempt_evaluation(
    response_text: str,
    tutor_state: Optional[dict],
    attempt_evaluation: object,
) -> tuple[str, Optional[dict]]:
    data = _attempt_evaluation_to_dict(attempt_evaluation)
    verdict = str(data.get("verdict") or "").strip().lower()
    if verdict not in {"correct", "partly_correct", "not_yet", "needs_attempt"}:
        return response_text, tutor_state

    state = dict(tutor_state or {})
    state["verdict"] = verdict
    try:
        existing_confidence = float(state.get("confidence") or 0.0)
    except (TypeError, ValueError):
        existing_confidence = 0.0
    try:
        graph_confidence = float(data.get("confidence") or 0.0)
    except (TypeError, ValueError):
        graph_confidence = 0.0
    state["confidence"] = round(max(0.0, min(1.0, max(existing_confidence, graph_confidence))), 2)
    if verdict == "correct":
        state["phase"] = "practice" if state.get("phase") == "check" else state.get("phase", "teach")
    if data.get("next_action"):
        state["next_action"] = _trim_tutor_text(data.get("next_action"), state.get("next_action") or "Try the next small step", 120)
    if data.get("expected_answer"):
        state["expected_step_answer"] = _trim_tutor_text(data.get("expected_answer"), state.get("expected_step_answer") or "", 300)
    if data.get("is_final_answer"):
        state["phase"] = "review" if data.get("final_answer_correct") else state.get("phase", "check")
        if data.get("final_answer_correct") and state.get("total_steps"):
            state["current_step"] = state.get("total_steps")
    if data.get("misconception") and verdict in {"partly_correct", "not_yet"}:
        misconceptions = _trim_string_list(state.get("misconceptions"), 12, 120)
        misconception = _trim_tutor_text(data.get("misconception"), "", 120)
        if misconception and misconception not in misconceptions:
            misconceptions.append(misconception)
        state["misconceptions"] = misconceptions[-12:]
    answer = response_text or ""
    if verdict == "correct" and not re.search(r"\b(correct|right|exactly|yes)\b", answer, flags=re.I):
        reason = _trim_tutor_text(data.get("rationale"), "that matches the requested step", 140)
        answer = f"Correct — {reason}.\n\n{answer}".strip()
    elif verdict in {"partly_correct", "not_yet"} and re.search(r"\b(correct|exactly)\b", answer, flags=re.I):
        answer = re.sub(r"\bCorrect\b\s*[—:-]?\s*", "", answer, count=1, flags=re.I).strip()

    return answer, state

def _tutor_state_row_to_payload(row: models.ChatTutorState | None) -> Optional[dict]:
    if not row:
        return None
    payload = _normalize_tutor_state({
        "level": row.level,
        "phase": row.phase,
        "verdict": row.verdict,
        "confidence": row.confidence,
        "objective": row.objective,
        "next_action": row.next_action,
        "hint_level": row.hint_level,
        "current_step": row.current_step,
        "total_steps": row.total_steps,
        "expected_step_answer": row.expected_step_answer,
        "final_answer": row.final_answer,
        "skills_used": row.skills_used,
        "misconceptions": row.misconceptions,
        "mastery_score": row.mastery_score,
        "correct_streak": row.correct_streak,
        "wrong_streak": row.wrong_streak,
        "lesson_plan": row.lesson_plan,
    }, reply_style=row.reply_style or "guided")
    payload["attempts"] = row.attempts or 0
    payload["correct_count"] = row.correct_count or 0
    return payload

def _get_tutor_session_state(db: Session, chat_id: Optional[int], user_id: int) -> Optional[dict]:
    if not chat_id:
        return None
    with db.no_autoflush:
        row = (
            db.query(models.ChatTutorState)
            .filter(
                models.ChatTutorState.chat_session_id == chat_id,
                models.ChatTutorState.user_id == user_id,
            )
            .first()
        )
    return _tutor_state_row_to_payload(row)

def _persist_tutor_session_state(
    db: Session,
    chat_id: Optional[int],
    user_id: int,
    tutor_state: Optional[dict],
    reply_style: str,
    tutor_options: Optional[list[dict]] = None,
    tutor_choice: Optional[str] = None,
) -> Optional[dict]:
    if not chat_id or not tutor_state:
        return tutor_state

    with db.no_autoflush:
        session_exists = (
            db.query(models.ChatSession.id)
            .filter(
                models.ChatSession.id == chat_id,
                models.ChatSession.user_id == user_id,
            )
            .first()
        )
        user_exists = db.query(models.User.id).filter(models.User.id == user_id).first()
        if not session_exists or not user_exists:
            logger.warning(
                "Skipping tutor state persistence for missing FK target: chat_id=%s user_id=%s",
                chat_id,
                user_id,
            )
            return tutor_state

        row = (
            db.query(models.ChatTutorState)
            .filter(
                models.ChatTutorState.chat_session_id == chat_id,
                models.ChatTutorState.user_id == user_id,
            )
            .first()
        )
    if not row:
        row = models.ChatTutorState(chat_session_id=chat_id, user_id=user_id)
        db.add(row)

    verdict = tutor_state.get("verdict")
    is_student_attempt = bool(tutor_choice) or verdict in {"correct", "partly_correct", "not_yet"}

    if is_student_attempt:
        row.attempts = (row.attempts or 0) + 1

    if tutor_choice:
        row.last_choice = tutor_choice

    if is_student_attempt and verdict == "correct":
        row.correct_count = (row.correct_count or 0) + 1

    row.level = tutor_state.get("level", "intermediate")
    row.phase = tutor_state.get("phase", "teach")
    row.verdict = tutor_state.get("verdict", "not_applicable")
    row.confidence = tutor_state.get("confidence", 0.65)
    row.objective = tutor_state.get("objective")
    row.next_action = tutor_state.get("next_action")
    row.hint_level = tutor_state.get("hint_level", 2)
    row.current_step = tutor_state.get("current_step", row.current_step or 1)
    row.total_steps = tutor_state.get("total_steps", row.total_steps or 0)
    row.expected_step_answer = tutor_state.get("expected_step_answer") or row.expected_step_answer
    row.final_answer = tutor_state.get("final_answer") or row.final_answer
    row.skills_used = tutor_state.get("skills_used") or row.skills_used or []
    row.misconceptions = tutor_state.get("misconceptions") or row.misconceptions or []
    row.mastery_score = tutor_state.get("mastery_score", row.mastery_score or 0.0)
    if not is_student_attempt:
        row.correct_streak = tutor_state.get("correct_streak", row.correct_streak or 0)
        row.wrong_streak = tutor_state.get("wrong_streak", row.wrong_streak or 0)
    row.lesson_plan = tutor_state.get("lesson_plan") or row.lesson_plan
    if is_student_attempt and verdict == "correct":
        row.correct_streak = (row.correct_streak or 0) + 1
        row.wrong_streak = 0
    elif is_student_attempt and verdict in {"partly_correct", "not_yet"}:
        row.wrong_streak = (row.wrong_streak or 0) + 1
        row.correct_streak = 0

    if row.correct_streak >= 2 and row.level == "beginner":
        row.level = "intermediate"
    elif row.correct_streak >= 2 and row.level == "intermediate":
        row.level = "advanced"
    elif row.wrong_streak >= 2 and row.level == "advanced":
        row.level = "intermediate"
    elif row.wrong_streak >= 2 and row.level == "intermediate":
        row.level = "beginner"

    total_attempts = max(1, row.attempts or 0)
    row.mastery_score = round(min(1.0, max(0.0, (row.correct_count or 0) / total_attempts)), 2)
    row.reply_style = reply_style or "guided"
    row.last_options = tutor_options or []
    row.updated_at = datetime.now(timezone.utc)
    recorded_weaknesses = _record_tutor_weakness_signals(db, user_id, tutor_state, tutor_choice)
    payload = _tutor_state_row_to_payload(row)
    if recorded_weaknesses:
        payload["auto_weakness_topics"] = recorded_weaknesses
    return payload

def _tutor_mode_prompt_prefix(reply_style: str = "guided") -> str:
    return f"{tutor_contract_instruction(reply_style)}\n\n"

def _assert_user_matches_request(user_id: Optional[str], current_user: models.User) -> None:
    if user_id is None:
        return
    requested = str(user_id).strip().lower()
    allowed = {
        (current_user.username or "").strip().lower(),
        (current_user.email or "").strip().lower(),
    }
    if requested and requested not in allowed:
        raise HTTPException(status_code=403, detail="Access denied")

def _resolve_chat_session_id(db: Session, chat_id: Optional[str], user_id: int) -> Optional[int]:
    if not chat_id:
        return None
    identifier = str(chat_id).strip()
    if identifier.isdigit():
        return int(identifier)
    session = db.query(models.ChatSession).filter(
        models.ChatSession.public_token == identifier,
        models.ChatSession.user_id == user_id,
    ).first()
    return session.id if session else None

def _context_only_fallback_answer(user_id: str, question: str, context_doc_ids: list[str], use_hs_context: bool = True) -> str:
    selected_ids = [str(d).strip() for d in (context_doc_ids or []) if str(d).strip()]
    if not selected_ids or not use_hs_context:
        return call_ai(question)

    try:
        from services import context_store
        if not context_store.available():
            return call_ai(question)

        q = (question or "").strip() or "Summarize the selected context."
        results = context_store.search_context(
            query=q,
            user_id=str(user_id),
            use_hs=False,
            top_k=8,
            doc_ids=selected_ids,
        )
        chunks = [r.get("text", "").strip() for r in results if r.get("text")]
        if not chunks:
            return (
                "I couldn’t find anything related to that in the selected context documents. "
                "Please select more relevant context pages/files and try again."
            )

        context_blob = "\n\n".join(f"--- Chunk {i+1} ---\n{c}" for i, c in enumerate(chunks[:6]))
        prompt = (
            "You are a tutor in STRICT CONTEXT-ONLY mode.\n"
            "Use only the provided context chunks.\n"
            "Do not use outside knowledge.\n"
            "If the answer is not supported by these chunks, say so clearly.\n\n"
            f"Question:\n{q}\n\n"
            f"Selected context chunks:\n{context_blob}\n"
        )
        return call_ai(prompt)
    except Exception as e:
        logger.warning(f"Context-only fallback failed: {e}")
        return (
            "I couldn't complete a context-only answer right now. "
            "Please try again in a moment."
        )

async def _run_chat_ml_pipeline(
    db: Session,
    user: models.User,
    question: str,
    chat_id_int: Optional[int],
    session_msg_count: int,
    tutor_mode: bool,
):
    """Runs MessageMLPipeline (BKT mastery update, StrategyBandit selection +
    reward-queueing, frustration/engagement tracking) and returns
    (ml_output, ml_addendum) for the caller to feed into tutor.invoke().

    Shared by /ask/ and /ask_simple/ so both actually exercise the bandit/BKT
    system -- /ask_simple/ (what the web frontend calls) used to skip this
    entirely and call the tutor graph directly, so StrategyBandit selection,
    real-time chat mastery updates, and frustration/engagement trend tracking
    never happened for web users even though the underlying pipeline itself
    was fully wired and tested (just never reached via that endpoint).
    """
    ml_output = None
    ml_addendum = ""
    try:
        from services.ml_pipeline import MessageMLPipeline, SessionContext
        from services.memory_service import get_memory_service

        session_state = None
        if chat_id_int:
            session_state = db.query(models.CerbylSessionState).filter_by(
                session_id=chat_id_int
            ).first()
            if not session_state:
                session_state = models.CerbylSessionState(
                    session_id=chat_id_int,
                    user_id=user.id,
                    started_at=datetime.now(timezone.utc),
                    message_count=0,
                )
                db.add(session_state)
                db.commit()
                db.refresh(session_state)

        ctx = SessionContext(
            session_id=chat_id_int,
            message_count=session_state.message_count if session_state else session_msg_count,
            current_concept_id=session_state.current_concept_id if session_state else None,
            messages_on_concept=(
                (session_state.messages_on_concept or {}).get(
                    session_state.current_concept_id or "", 0
                ) if session_state else 0
            ),
            frustration_trend=session_state.frustration_trend or [] if session_state else [],
            engagement_trend=session_state.engagement_trend or [] if session_state else [],
        )

        pipeline = MessageMLPipeline(None, get_memory_service())
        ml_output = await pipeline.process(question, str(user.id), ctx, db)
        weak_profile = pipeline.build_weak_concept_profile(db, user.id)
        ml_addendum = pipeline.build_system_prompt_addendum(
            ml_output,
            profile=weak_profile,
            require_follow_up=bool(tutor_mode),
        )

        if session_state and ml_output:
            session_state.message_count += 1
            session_state.last_message_at = datetime.now(timezone.utc)
            trend = session_state.frustration_trend or []
            trend.append(round(ml_output.frustration_score, 3))
            session_state.frustration_trend = trend[-10:]
            if ml_output.detected_concepts:
                session_state.current_concept_id = ml_output.detected_concepts[0]
                mon = session_state.messages_on_concept or {}
                cid = ml_output.detected_concepts[0]
                mon[cid] = mon.get(cid, 0) + 1
                session_state.messages_on_concept = mon
            db.commit()

    except Exception as _ml_err:
        logger.warning(f"[CHAT] ML pipeline skipped: {_ml_err}", exc_info=True)

    try:
        from services.context_agent import get_context_agent, LearningEvent

        _agent = get_context_agent()
        if _agent and ml_output:
            _event = LearningEvent(
                student_id=str(user.id),
                source="chat",
                event_type="message",
                concept_id=(ml_output.detected_concepts[0] if ml_output.detected_concepts else ""),
                concept_name="",
                session_id=chat_id_int,
                frustration=ml_output.frustration_score,
                intent=ml_output.intent,
                message=question[:300],
            )
            _agent.record_event(db, _event)
    except Exception as _ag_err:
        logger.debug(f"[CHAT] context agent skipped: {_ag_err}")

    return ml_output, ml_addendum

class ChatSessionCreate(BaseModel):
    user_id: str
    title: str = "New Chat"

class ChatMessageSave(BaseModel):
    chat_id: int
    user_message: str
    ai_response: str

class GenerateChatTitleRequest(BaseModel):
    chat_id: str
    user_id: str

class ChatFolderCreate(BaseModel):
    user_id: str
    name: str
    color: Optional[str] = "#D7B38C"
    parent_id: Optional[int] = None

class ChatUpdateFolder(BaseModel):
    chat_id: str
    folder_id: Optional[int] = None

@router.post("/ask/")
async def ask_ai(
    user_id: str = Form(...),
    question: str = Form(...),
    chat_id: Optional[str] = Form(None),
    use_hs_context: bool = Form(True),
    context_doc_ids: Optional[str] = Form(None),
    tutor_mode: bool = Form(False),
    tutor_reply_style: str = Form("guided"),
    tutor_choice: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    try:
        selected_doc_ids = [x.strip() for x in (context_doc_ids or "").split(",") if x.strip()][:200]
        chat_id_int = _resolve_chat_session_id(db, chat_id, current_user.id)

        _assert_user_matches_request(user_id, current_user)
        user = current_user

        logger.info(
            f"[CHAT ROUTE] message | user={user.id} "
            f"HS_MODE={'ON  <-- curriculum RAG will run' if use_hs_context else 'OFF <-- no RAG, model-only'} "
            f"| q='{question[:80]}'"
        )

        today = datetime.now(timezone.utc).date()
        daily_metric = db.query(models.DailyLearningMetrics).filter(
            and_(
                models.DailyLearningMetrics.user_id == user.id,
                models.DailyLearningMetrics.date == today,
            )
        ).first()

        if not daily_metric:
            daily_metric = models.DailyLearningMetrics(
                user_id=user.id,
                date=today,
                questions_answered=0,
                correct_answers=0,
                sessions_completed=0,
                time_spent_minutes=0.0,
                accuracy_rate=0.0,
                engagement_score=0.0,
            )
            db.add(daily_metric)
            db.commit()
            db.refresh(daily_metric)

        if chat_id_int:
            chat_session = db.query(models.ChatSession).filter(
                models.ChatSession.id == chat_id_int,
                models.ChatSession.user_id == user.id,
            ).first()
            if not chat_session:
                raise HTTPException(status_code=404, detail="Chat session not found")

        chat_history_for_tutor = []
        if chat_id_int:
            recent_msgs_for_ctx = (
                db.query(models.ChatMessage)
                .filter(models.ChatMessage.chat_session_id == chat_id_int)
                .order_by(models.ChatMessage.timestamp.desc())
                .limit(20)
                .all()
            )
            chat_history_for_tutor = [
                {"user": msg.user_message, "ai": msg.ai_response}
                for msg in reversed(recent_msgs_for_ctx)
            ]
        tutor_session_state = _get_tutor_session_state(db, chat_id_int, user.id) if tutor_mode else None

        ml_output, ml_addendum = await _run_chat_ml_pipeline(
            db, user, question, chat_id_int, len(chat_history_for_tutor), tutor_mode,
        )

        from tutor.graph import get_tutor

        tutor = get_tutor()
        if tutor:
            result = await tutor.invoke(
                user_id=str(user.id),
                user_input=question,
                chat_id=chat_id_int,
                chat_history=chat_history_for_tutor,
                use_hs_context=bool(use_hs_context),
                context_doc_ids=selected_doc_ids,
                context_only=bool(selected_doc_ids),
                ml_addendum=ml_addendum,
                tutor_mode=bool(tutor_mode),
                tutor_reply_style=tutor_reply_style,
                tutor_choice=tutor_choice,
                tutor_session_state=tutor_session_state,
            )
            response_text = result.get("response", "")
            if tutor_mode:
                response_text, tutor_options, tutor_state = _parse_tutor_response(
                    response_text,
                    tutor_reply_style,
                    tutor_choice,
                )
                tutor_state = _apply_tutor_plan(tutor_state, result.get("tutor_plan"))
                response_text, tutor_state = _apply_attempt_evaluation(
                    response_text,
                    tutor_state,
                    result.get("attempt_evaluation"),
                )
            else:
                tutor_options = []
                tutor_state = None
            try:
                from services.math_processor import process_math_in_response
                response_text = process_math_in_response(response_text)
            except Exception:
                pass
        else:
            response_text = _context_only_fallback_answer(str(user.id), question, selected_doc_ids, use_hs_context)
            tutor_options = []
            tutor_state = _normalize_tutor_state({}, tutor_reply_style, tutor_choice) if tutor_mode else None

        if ml_output:
            try:
                ml_log = models.MessageMLLog(
                    session_id=chat_id_int,
                    user_id=user.id,
                    message_text=question[:500],
                    intent_class=ml_output.intent,
                    concept_ids=ml_output.detected_concepts,
                    frustration_score=ml_output.frustration_score,
                    engagement_score=ml_output.engagement_score,
                    cognitive_state=ml_output.cognitive_state,
                    archetype=ml_output.archetype,
                    response_strategy=ml_output.response_strategy,
                    kt_delta=ml_output.kt_after,
                    memories_used=ml_output.memories_used,
                    messages_this_session=len(chat_history_for_tutor) + 1,
                )
                db.add(ml_log)
            except Exception:
                pass

        if chat_id_int:
            msg = models.ChatMessage(
                chat_session_id=chat_id_int,
                user_id=user.id,
                user_message=question,
                ai_response=response_text,
                is_user=True,
                timestamp=datetime.now(timezone.utc),
            )
            db.add(msg)

            session = db.query(models.ChatSession).filter(
                models.ChatSession.id == chat_id_int,
                models.ChatSession.user_id == user.id,
            ).first()
            if session:
                session.updated_at = datetime.now(timezone.utc)
                # NOTE: `user_question`/`model_question` are not defined anywhere in this
                # function (they belong to a different handler further down in this same
                # file) -- this was an unconditional NameError on every turn that had a
                # chat_id, silently swallowed by the blanket `except Exception` below and
                # masked as a generic "I encountered an error" reply. `question` is this
                # function's actual Form parameter holding the user's message.
                _apply_fallback_chat_title(session, question)

            if tutor_mode:
                tutor_state = _persist_tutor_session_state(
                    db,
                    chat_id_int,
                    user.id,
                    tutor_state,
                    tutor_reply_style,
                    tutor_options,
                    tutor_choice,
                )

            try:
                from services.gamification_system import award_points
                award_points(db, user.id, "ai_chat", {"question": question})
            except Exception:
                pass

            db.commit()

        _intent_result = None
        _computed_confidence = 0.72
        try:
            from services.intent_engine import CerbylIntentEngine
            _engine = CerbylIntentEngine.get()
            _intent_result = _engine.classify(question, user.id)
            _computed_confidence = _engine.estimate_response_confidence(
                result=_intent_result,
                response_text=response_text,
                engagement_score=ml_output.engagement_score if ml_output else 0.5,
                frustration_score=ml_output.frustration_score if ml_output else 0.0,
                p_mastery=ml_output.p_mastery if ml_output else 0.1,
            )
            _implicit_label = (
                "INSTRUCTION" if _intent_result.is_instruction() else
                "CASUAL"      if _intent_result.is_casual()      else
                _intent_result.label
            )
            _implicit_weight = 1.5 if _implicit_label == "INSTRUCTION" else 0.7
            _engine.record_signal(question, _implicit_label, weight=_implicit_weight)
        except Exception as _ie_err:
            logger.debug("[CHAT] IntentEngine skipped: %s", _ie_err)

        return {
            "answer":            response_text,
            "ai_confidence":     _computed_confidence,
            "topics_discussed":  ml_output.detected_concepts if ml_output else [],
            "query_type":        ml_output.intent if ml_output else "conversational_learning",
            "intent_class":      _intent_result.label if _intent_result else "LEARN_CONCEPT",
            "intent_proba":      _intent_result.proba if _intent_result else {},
            "active_rules":      [{"domain": r.domain, "negated": r.negated}
                                  for r in (_intent_result.active_rules if _intent_result else [])],
            "questions_today":   daily_metric.questions_answered,
            "frustration_score": ml_output.frustration_score if ml_output else 0.0,
            "response_strategy": ml_output.response_strategy if ml_output else "",
            "tutor_mode": bool(tutor_mode),
            "tutor_reply_style": tutor_reply_style,
            "tutor_options": tutor_options,
            "tutor_state": tutor_state,
        }

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Error in /api/ask/: {e}", exc_info=True)
        return {
            "answer": "I apologize, but I encountered an error. Could you please rephrase your question?",
            "ai_confidence": 0.3,
            "topics_discussed": ["error"],
            "query_type": "error",
        }

@router.post("/ask_simple/")
async def ask_simple(
    user_id: str = Form(...),
    question: str = Form(...),
    original_question: Optional[str] = Form(None),
    chat_id: Optional[str] = Form(None),
    use_hs_context: bool = Form(True),
    context_doc_ids: Optional[str] = Form(None),
    tutor_mode: bool = Form(False),
    tutor_reply_style: str = Form("guided"),
    tutor_choice: Optional[str] = Form(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    try:
        selected_doc_ids = [x.strip() for x in (context_doc_ids or "").split(",") if x.strip()][:200]
        _assert_user_matches_request(user_id, current_user)
        user = current_user
        user_question = (original_question or question or "").strip()
        model_question = question
        tutor_user_question = _strip_internal_graph_guidance(user_question or model_question)
        effective_tutor_input = tutor_user_question if tutor_mode else model_question

        chat_id_int = _resolve_chat_session_id(db, chat_id, current_user.id)
        if _load_test_fallback_enabled(user):
            response_text = _load_test_answer(effective_tutor_input)
            return {
                "answer": response_text,
                "ai_confidence": 0.45,
                "intent_class": "LEARN_CONCEPT",
                "active_rules": [],
                "topics_discussed": [],
                "query_type": "load_test_fallback",
                "tutor_mode": bool(tutor_mode),
                "tutor_reply_style": tutor_reply_style,
                "tutor_options": [],
                "tutor_state": None,
            }

        chat_history = []
        if chat_id_int:
            chat_session = db.query(models.ChatSession).filter(
                models.ChatSession.id == chat_id_int,
                models.ChatSession.user_id == user.id,
            ).first()
            if not chat_session:
                raise HTTPException(status_code=404, detail="Chat session not found")

            recent_msgs = (
                db.query(models.ChatMessage)
                .filter(models.ChatMessage.chat_session_id == chat_id_int)
                .order_by(models.ChatMessage.timestamp.desc())
                .limit(20)
                .all()
            )
            chat_history = [
                {"user": msg.user_message, "ai": msg.ai_response}
                for msg in reversed(recent_msgs)
            ]
        tutor_session_state = _get_tutor_session_state(db, chat_id_int, user.id) if tutor_mode else None

        ml_output, ml_addendum = await _run_chat_ml_pipeline(
            db, user, user_question or model_question, chat_id_int, len(chat_history), tutor_mode,
        )

        from tutor.graph import get_tutor

        tutor = get_tutor()
        if tutor:
            result = await tutor.invoke(
                user_id=str(user.id),
                user_input=effective_tutor_input,
                chat_id=chat_id_int,
                chat_history=chat_history,
                use_hs_context=bool(use_hs_context),
                context_doc_ids=selected_doc_ids,
                context_only=bool(selected_doc_ids),
                ml_addendum=ml_addendum,
                tutor_mode=bool(tutor_mode),
                tutor_reply_style=tutor_reply_style,
                tutor_choice=tutor_choice,
                tutor_session_state=tutor_session_state,
            )
            response_text = result.get("response", "")
            if tutor_mode:
                response_text, tutor_options, tutor_state = _parse_tutor_response(
                    response_text,
                    tutor_reply_style,
                    tutor_choice,
                )
                tutor_state = _apply_tutor_plan(tutor_state, result.get("tutor_plan"))
                response_text, tutor_state = _apply_attempt_evaluation(
                    response_text,
                    tutor_state,
                    result.get("attempt_evaluation"),
                )
            else:
                tutor_options = []
                tutor_state = None
            try:
                from services.math_processor import process_math_in_response
                response_text = process_math_in_response(response_text)
            except Exception:
                pass
        else:
            response_text = _context_only_fallback_answer(str(user.id), effective_tutor_input, selected_doc_ids, use_hs_context)
            tutor_options = []
            tutor_state = _normalize_tutor_state({}, tutor_reply_style, tutor_choice) if tutor_mode else None

        if chat_id_int:
            msg = models.ChatMessage(
                chat_session_id=chat_id_int,
                user_id=user.id,
                user_message=user_question or model_question,
                ai_response=response_text,
                timestamp=datetime.now(timezone.utc),
            )
            db.add(msg)

            session = db.query(models.ChatSession).filter(
                models.ChatSession.id == chat_id_int,
                models.ChatSession.user_id == user.id,
            ).first()
            if session:
                session.updated_at = datetime.now(timezone.utc)
                _apply_fallback_chat_title(
                    session,
                    user_question or model_question or "Uploaded files",
                )

            if tutor_mode:
                tutor_state = _persist_tutor_session_state(
                    db,
                    chat_id_int,
                    user.id,
                    tutor_state,
                    tutor_reply_style,
                    tutor_options,
                    tutor_choice,
                )

            try:
                from services.gamification_system import award_points
                award_points(db, user.id, "ai_chat")
            except Exception:
                pass

            db.commit()

        _intent_result = None
        _computed_confidence = 0.72
        try:
            from services.intent_engine import CerbylIntentEngine
            _engine = CerbylIntentEngine.get()
            _intent_result = _engine.classify(user_question or model_question, user.id)
            _computed_confidence = _engine.estimate_response_confidence(
                result=_intent_result,
                response_text=response_text,
            )
            _engine.record_signal(
                user_question or model_question,
                "INSTRUCTION" if _intent_result.is_instruction() else _intent_result.label,
                weight=1.5 if _intent_result.is_instruction() else 0.7,
            )
        except Exception as _ie_err:
            logger.debug("[CHAT/simple] IntentEngine skipped: %s", _ie_err)

        return {
            "answer":        response_text,
            "ai_confidence": _computed_confidence,
            "intent_class":  _intent_result.label if _intent_result else "LEARN_CONCEPT",
            "active_rules":  [{"domain": r.domain, "negated": r.negated}
                              for r in (_intent_result.active_rules if _intent_result else [])],
            "topics_discussed":  ml_output.detected_concepts if ml_output else [],
            "frustration_score": ml_output.frustration_score if ml_output else 0.0,
            "response_strategy": ml_output.response_strategy if ml_output else "",
            "query_type":    "conversational_learning",
            "tutor_mode": bool(tutor_mode),
            "tutor_reply_style": tutor_reply_style,
            "tutor_options": tutor_options,
            "tutor_state": tutor_state,
        }

    except HTTPException:
        raise
    except Exception as e:
        if _is_usage_limit_error(e):
            fallback_answer = _load_test_answer((original_question or question or "").strip())
            return {
                "answer": fallback_answer,
                "ai_confidence": 0.45,
                "topics_discussed": [],
                "query_type": "provider_quota_fallback",
                "tutor_mode": bool(tutor_mode),
                "tutor_reply_style": tutor_reply_style,
                "tutor_options": [],
                "tutor_state": None,
            }
        logger.error(f"Error in /api/ask_simple/: {e}", exc_info=True)
        return {
            "answer": "I encountered an error processing your request.",
            "ai_confidence": 0.3,
            "topics_discussed": ["error"],
            "query_type": "error",
        }

@router.post("/ask_with_files/")
async def ask_with_files(
    user_id: str = Form(...),
    question: str = Form(...),
    original_question: Optional[str] = Form(None),
    chat_id: Optional[str] = Form(None),
    use_hs_context: bool = Form(True),
    context_doc_ids: Optional[str] = Form(None),
    tutor_mode: bool = Form(False),
    tutor_reply_style: str = Form("guided"),
    tutor_choice: Optional[str] = Form(None),
    files: List[UploadFile] = File(default=[]),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    try:
        selected_doc_ids = [x.strip() for x in (context_doc_ids or "").split(",") if x.strip()][:200]
        _assert_user_matches_request(user_id, current_user)
        user = current_user
        user_question = (original_question or question or "").strip()
        model_question = question
        tutor_user_question = _strip_internal_graph_guidance(user_question or model_question)

        chat_id_int = _resolve_chat_session_id(db, chat_id, current_user.id)
        if chat_id_int:
            chat_session = db.query(models.ChatSession).filter(
                models.ChatSession.id == chat_id_int,
                models.ChatSession.user_id == user.id,
            ).first()
            if not chat_session:
                raise HTTPException(status_code=404, detail="Chat session not found")

        image_payloads: list[dict] = []
        text_extracts: list[str] = []
        extraction_errors: list[str] = []
        saved_metadata: list[dict] = []
        storage = StorageService.get_storage()
        storage_type = getattr(storage, "storage_type", "local")

        if not files:
            raise HTTPException(
                status_code=422,
                detail="No attachment reached the server. Please reattach the file and try again.",
            )

        unsupported_files: list[str] = []
        for upload in files[:_MAX_IMAGES_PER_MESSAGE]:
            raw = await upload.read()
            if not raw:
                unsupported_files.append(f"{upload.filename or 'upload'}: empty file")
                continue

            fname = upload.filename or "upload"
            lower_fname = fname.lower()
            mime = _normalized_upload_mime(upload.content_type, fname)

            if mime in _SUPPORTED_IMAGE_TYPES:
                if len(raw) > _MAX_IMAGE_BYTES:
                    unsupported_files.append(f"{fname}: image exceeds the 20 MB limit")
                    continue
                safe_name = _safe_filename(user.id, fname)
                storage_path = _store_chat_upload(storage, raw, f"chat_images/{user.id}/{safe_name}", safe_name, mime)
                image_payloads.append({"data": raw, "mime_type": mime, "filename": fname})
                saved_metadata.append({
                    "filename": fname,
                    "mime_type": mime,
                    "size": len(raw),
                    "storage_path": storage_path,
                    "is_image": True,
                })

            elif mime == "application/pdf" or fname.lower().endswith(".pdf"):
                safe_name = _safe_filename(user.id, fname)
                storage_path = _store_chat_upload(storage, raw, f"chat_files/{user.id}/{safe_name}", safe_name, "application/pdf")
                analysis_bytes = _read_uploaded_bytes(storage, storage_path, raw)
                extraction = _extract_pdf_text_details(analysis_bytes)
                extracted = (extraction.get("text") or "").strip()
                if extracted:
                    parser = extraction.get("parser") or "pdf"
                    pages = extraction.get("page_count") or 0
                    text_extracts.append(f"[PDF: {fname} | parser: {parser} | pages: {pages}]\n{extracted[:12000]}")
                else:
                    warnings = "; ".join(extraction.get("warnings") or [])
                    extraction_errors.append(
                        f"{fname}: no readable PDF text was extracted"
                        + (f" ({warnings})" if warnings else "")
                    )
                saved_metadata.append({
                    "filename": fname,
                    "mime_type": "application/pdf",
                    "size": len(raw),
                    "storage_path": storage_path,
                    "is_image": False,
                })

            elif mime in _SUPPORTED_DOCX_TYPES or lower_fname.endswith(".docx"):
                safe_name = _safe_filename(user.id, fname)
                content_type = mime or "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                storage_path = _store_chat_upload(storage, raw, f"chat_files/{user.id}/{safe_name}", safe_name, content_type)
                analysis_bytes = _read_uploaded_bytes(storage, storage_path, raw)
                extracted = _extract_docx_text(analysis_bytes)
                if extracted:
                    text_extracts.append(f"[DOCX: {fname}]\n{extracted[:12000]}")
                else:
                    extraction_errors.append(f"{fname}: no readable DOCX text was extracted")
                saved_metadata.append({
                    "filename": fname,
                    "mime_type": content_type,
                    "size": len(raw),
                    "storage_path": storage_path,
                    "is_image": False,
                })
            elif mime in _SUPPORTED_TEXT_TYPES or lower_fname.endswith(_SUPPORTED_TEXT_EXTENSIONS):
                safe_name = _safe_filename(user.id, fname)
                content_type = mime or "text/plain"
                storage_path = _store_chat_upload(storage, raw, f"chat_files/{user.id}/{safe_name}", safe_name, content_type)
                analysis_bytes = _read_uploaded_bytes(storage, storage_path, raw)
                extracted = analysis_bytes.decode("utf-8", errors="replace").strip()
                if extracted:
                    text_extracts.append(f"[File: {fname}]\n{extracted[:12000]}")
                else:
                    extraction_errors.append(f"{fname}: no readable text was extracted")
                saved_metadata.append({
                    "filename": fname,
                    "mime_type": content_type,
                    "size": len(raw),
                    "storage_path": storage_path,
                    "is_image": False,
                })
            else:
                unsupported_files.append(f"{fname}: unsupported file type")

        if extraction_errors and not text_extracts and not image_payloads:
            raise HTTPException(
                status_code=422,
                detail=(
                    "I could not extract readable content from the uploaded file(s): "
                    + "; ".join(extraction_errors)
                    + ". If this is a scanned PDF, upload a text-based PDF or run OCR first."
                ),
            )
        if not saved_metadata:
            detail = "; ".join(unsupported_files) or "no supported file content was received"
            raise HTTPException(
                status_code=422,
                detail=f"The attachment could not be processed: {detail}. Please reattach a supported file.",
            )

        enriched_question = model_question.strip() or "Please analyze the attached content."
        if tutor_mode:
            enriched_question = _tutor_mode_prompt_prefix(tutor_reply_style) + enriched_question
        if text_extracts:
            enriched_question += "\n\n" + "\n\n".join(text_extracts)
        if extraction_errors:
            enriched_question += "\n\n[Attachment extraction warnings]\n" + "\n".join(extraction_errors)
        if image_payloads:
            img_count = len(image_payloads)
            enriched_question = (
                f"{enriched_question}\n\n"
                f"[{img_count} image{'s' if img_count > 1 else ''} attached — analyse and respond]"
            )

        chat_history = []
        if chat_id_int:
            recent = (
                db.query(models.ChatMessage)
                .filter(models.ChatMessage.chat_session_id == chat_id_int)
                .order_by(models.ChatMessage.timestamp.desc())
                .limit(20)
                .all()
            )
            chat_history = [
                {"user": m.user_message, "ai": m.ai_response}
                for m in reversed(recent)
            ]
        tutor_session_state = _get_tutor_session_state(db, chat_id_int, user.id) if tutor_mode else None

        response_text = ""
        ai_provider = "AI"
        vision_error = ""

        from deps import unified_ai as _unified_ai
        from services.ai_utils import NoVisionProviderError
        ai_client = _unified_ai

        if image_payloads and _is_person_identity_request(user_question or model_question):
            response_text = (
                "I can’t identify or confirm a person’s identity from an image. "
                "I can describe their visible appearance, clothing, setting, or actions instead."
            )
            ai_provider = "privacy_guard"

        # A directly attached image belongs to the current message. Selected
        # Vault/context documents must not suppress multimodal analysis.
        if ai_client and image_payloads and not response_text:
            try:
                response_text = (
                    await run_in_threadpool(
                        ai_client.generate_with_images,
                        enriched_question,
                        image_payloads,
                        2000,
                        0.7,
                    )
                    or ""
                ).strip()
                ai_provider = "vision"
            except NoVisionProviderError as exc:
                vision_error = str(exc) or "No vision-capable AI provider is configured"
                logger.warning(
                    "No vision provider configured; image response will not fall back to chat history"
                )
            except Exception as e:
                _raise_if_usage_limit_error(e)
                vision_error = f"Image analysis failed: {e}"
                logger.warning(
                    "Vision call failed (%s); image response will not fall back to chat history",
                    e,
                )

        if image_payloads and not response_text:
            return {
                "answer": "",
                "attachment_error": (
                    "I received the image, but could not analyze it right now. "
                    "Please retry in a moment."
                ),
                "attachment_error_detail": vision_error or "Vision service returned an empty response",
                "ai_confidence": 0.0,
                "topics_discussed": [],
                "query_type": "multimodal_error",
                "files_processed": len(saved_metadata),
                "file_summaries": [
                    {"file_name": m["filename"], "is_image": m["is_image"], "size": m["size"]}
                    for m in saved_metadata
                ],
                "has_file_context": True,
                "images_analyzed": 0,
                "ai_provider": "vision_error",
                "storage_type": storage_type,
                "tutor_mode": bool(tutor_mode),
                "tutor_reply_style": tutor_reply_style,
                "tutor_options": [],
                "tutor_state": None,
            }

        if not response_text:
            tutor_input = (tutor_user_question if tutor_mode else model_question).strip() or "What can you help me with?"
            if text_extracts:
                tutor_input += "\n\n" + "\n\n".join(text_extracts)
            if extraction_errors:
                tutor_input += "\n\n[Attachment extraction warnings]\n" + "\n".join(extraction_errors)
            try:
                from tutor.graph import get_tutor
                tutor = get_tutor()
                if tutor:
                    result = await tutor.invoke(
                        user_id=str(user.id),
                        user_input=tutor_input,
                        chat_id=chat_id_int,
                        chat_history=chat_history,
                        use_hs_context=bool(use_hs_context),
                        context_doc_ids=selected_doc_ids,
                        context_only=bool(selected_doc_ids),
                        tutor_mode=bool(tutor_mode),
                        tutor_reply_style=tutor_reply_style,
                        tutor_choice=tutor_choice,
                        tutor_session_state=tutor_session_state,
                    )
                    response_text = result.get("response", "")
                else:
                    response_text = _context_only_fallback_answer(str(user.id), tutor_input, selected_doc_ids, use_hs_context)
            except Exception as e:
                _raise_if_usage_limit_error(e)
                logger.error(f"Tutor graph failed in ask_with_files: {e}")
                response_text = _context_only_fallback_answer(str(user.id), tutor_input, selected_doc_ids, use_hs_context)

        if tutor_mode:
            response_text, tutor_options, tutor_state = _parse_tutor_response(
                response_text,
                tutor_reply_style,
                tutor_choice,
            )
            tutor_state = _apply_tutor_plan(
                tutor_state,
                result.get("tutor_plan") if "result" in locals() else None,
            )
            response_text, tutor_state = _apply_attempt_evaluation(
                response_text,
                tutor_state,
                result.get("attempt_evaluation") if "result" in locals() else None,
            )
        else:
            tutor_options = []
            tutor_state = None
        try:
            from services.math_processor import process_math_in_response
            response_text = process_math_in_response(response_text)
        except Exception:
            pass

        if chat_id_int:
            msg = models.ChatMessage(
                chat_session_id=chat_id_int,
                user_id=user.id,
                user_message=user_question or model_question or "[image upload]",
                ai_response=response_text,
                timestamp=datetime.now(timezone.utc),
                image_metadata=json.dumps(saved_metadata) if saved_metadata else None,
            )
            db.add(msg)

            session = db.query(models.ChatSession).filter(
                models.ChatSession.id == chat_id_int,
                models.ChatSession.user_id == user.id,
            ).first()
            if session:
                session.updated_at = datetime.now(timezone.utc)
                _apply_fallback_chat_title(
                    session,
                    user_question or model_question or "Uploaded files",
                )

            if tutor_mode:
                tutor_state = _persist_tutor_session_state(
                    db,
                    chat_id_int,
                    user.id,
                    tutor_state,
                    tutor_reply_style,
                    tutor_options,
                    tutor_choice,
                )

            try:
                from services.gamification_system import award_points
                award_points(db, user.id, "ai_chat")
            except Exception:
                pass

            db.commit()

        file_summaries = [
            {"file_name": m["filename"], "is_image": m["is_image"], "size": m["size"]}
            for m in saved_metadata
        ]

        _intent_result = None
        _computed_confidence = 0.72
        try:
            from services.intent_engine import CerbylIntentEngine
            _engine = CerbylIntentEngine.get()
            _intent_result = _engine.classify(user_question or model_question, user.id)
            _computed_confidence = _engine.estimate_response_confidence(
                result=_intent_result,
                response_text=response_text,
            )
        except Exception as _ie_err:
            logger.debug("[CHAT/files] IntentEngine skipped: %s", _ie_err)

        return {
            "answer":          response_text,
            "ai_confidence":   _computed_confidence,
            "intent_class":    _intent_result.label if _intent_result else "LEARN_CONCEPT",
            "topics_discussed": [],
            "query_type":      "multimodal",
            "files_processed": len(saved_metadata),
            "file_summaries":  file_summaries,
            "has_file_context": bool(saved_metadata),
            "images_analyzed": len(image_payloads) if ai_provider == "vision" else 0,
            "ai_provider":     ai_provider,
            "storage_type":     storage_type,
            "tutor_mode":      bool(tutor_mode),
            "tutor_reply_style": tutor_reply_style,
            "tutor_options":   tutor_options,
            "tutor_state":     tutor_state,
        }

    except HTTPException:
        raise
    except Exception as e:
        _raise_if_usage_limit_error(e)
        logger.error(f"Error in /api/ask_with_files/: {e}", exc_info=True)
        return {
            "answer": "I encountered an error processing your files. Please try again.",
            "ai_confidence": 0.3,
            "topics_discussed": ["error"],
            "query_type": "error",
            "files_processed": 0,
        }

def _safe_filename(user_id: int, original: str) -> str:
    ts = int(datetime.now(timezone.utc).timestamp() * 1000)
    clean = re.sub(r"[^\w.\-]", "_", original)
    return f"{user_id}_{ts}_{clean}"

def _store_chat_upload(storage, raw: bytes, storage_key: str, safe_name: str, content_type: str) -> str:
    if getattr(storage, "storage_type", "local") == "local":
        dest = CHAT_UPLOAD_DIR / safe_name
        dest.write_bytes(raw)
        return str(dest)

    storage.upload_bytes(raw, storage_key, content_type)
    if hasattr(storage, "uri_for_path"):
        return storage.uri_for_path(storage_key)
    return storage_key

def _storage_key_from_uri(storage_path: str) -> str:
    parsed = urlparse(storage_path or "")
    if parsed.scheme in {"s3", "r2"}:
        return parsed.path.lstrip("/")
    return storage_path

def _read_uploaded_bytes(storage, storage_path: str, fallback: bytes) -> bytes:
    if getattr(storage, "storage_type", "local") == "local":
        return fallback
    if not hasattr(storage, "download_bytes"):
        raise HTTPException(status_code=500, detail="Configured storage does not support file readback")
    try:
        return storage.download_bytes(_storage_key_from_uri(storage_path))
    except Exception as exc:
        logger.error("Failed to read uploaded chat file from storage: %s", exc, exc_info=True)
        raise HTTPException(status_code=502, detail="Uploaded file was stored but could not be read back from S3")

def _extract_pdf_text_details(raw: bytes) -> dict:
    try:
        return extract_text_from_pdf_detailed(raw)
    except Exception as exc:
        logger.warning("PDF extraction failed in chat upload: %s", exc)
        return {
            "text": "",
            "parser": "",
            "page_count": 0,
            "non_empty_pages": 0,
            "warnings": [str(exc)],
        }

def _extract_docx_text(raw: bytes) -> str:
    try:
        from docx import Document
        doc = Document(io.BytesIO(raw))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text and p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n\n".join(paragraphs)
    except Exception:
        return ""

@router.post("/test_ai_simple")
async def test_ai_simple(
    question: str = Form(...),
    current_user: models.User = Depends(get_current_user),
):
    try:
        response = await call_ai_async(f"Answer this question in one sentence: {question}", max_tokens=200)
        return {"answer": response, "status": "success"}
    except Exception:
        return {"answer": "Error generating response", "status": "error"}

@router.post("/create_chat_session")
def create_chat_session(
    session_data: ChatSessionCreate,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    chat_session = models.ChatSession(user_id=current_user.id, title=session_data.title)
    db.add(chat_session)
    db.commit()
    db.refresh(chat_session)

    return {
        "id": chat_session.id,
        "uid": chat_session.public_token,
        "session_id": chat_session.id,
        "title": chat_session.title,
        "folder_id": chat_session.folder_id,
        "created_at": chat_session.created_at.isoformat() + "Z",
        "updated_at": chat_session.updated_at.isoformat() + "Z",
        "status": "success",
    }

@router.put("/rename_chat_session")
def rename_chat_session(
    data: dict,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    chat_id = data.get("chat_id")
    new_title = data.get("new_title")
    if not chat_id or not new_title:
        raise HTTPException(status_code=400, detail="chat_id and new_title are required")

    chat_session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        chat_id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    chat_session.title = new_title
    chat_session.updated_at = datetime.now(timezone.utc)
    db.commit()

    return {"status": "success", "chat_id": chat_id, "new_title": new_title}

@router.get("/get_chat_sessions")
def get_chat_sessions(
    limit: Optional[int] = Query(None, ge=1, le=500),
    offset: int = Query(0, ge=0),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    query = (
        db.query(models.ChatSession)
        .filter(models.ChatSession.user_id == current_user.id)
        .order_by(models.ChatSession.updated_at.desc())
    )
    if offset:
        query = query.offset(offset)
    if limit:
        query = query.limit(limit)
    sessions = query.all()

    return {
        "sessions": [
            {
                "id": s.id,
                "uid": s.public_token,
                "title": s.title,
                "folder_id": s.folder_id,
                "created_at": s.created_at.isoformat() + "Z" if s.created_at else None,
                "updated_at": s.updated_at.isoformat() + "Z" if s.updated_at else None,
            }
            for s in sessions
        ]
    }

@router.get("/get_chat_messages")
def get_chat_messages(
    chat_id: str = Query(...),
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        chat_id,
    ).first()
    if not session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    messages = (
        db.query(models.ChatMessage)
        .filter(models.ChatMessage.chat_session_id == session.id)
        .order_by(models.ChatMessage.timestamp.asc())
        .all()
    )
    tutor_state_row = (
        db.query(models.ChatTutorState)
        .filter(
            models.ChatTutorState.chat_session_id == session.id,
            models.ChatTutorState.user_id == current_user.id,
        )
        .first()
    )
    tutor_state_payload = _tutor_state_row_to_payload(tutor_state_row)

    result = []
    for msg in messages:
        result.append({
            "id": f"user_{msg.id}",
            "type": "user",
            "content": msg.user_message,
            "timestamp": msg.timestamp.isoformat() + "Z",
        })
        result.append({
            "id": f"ai_{msg.id}",
            "type": "ai",
            "content": msg.ai_response,
            "timestamp": msg.timestamp.isoformat() + "Z",
            "aiConfidence": 0.85,
        })
    if tutor_state_payload:
        for entry in reversed(result):
            if entry.get("type") == "ai":
                entry["tutorMode"] = True
                entry["tutorReplyMode"] = tutor_state_row.reply_style or "guided"
                entry["tutorState"] = tutor_state_payload
                entry["tutorOptions"] = tutor_state_row.last_options or []
                break
    return result

@router.get("/get_chat_history/{session_id}")
async def get_chat_history(
    session_id: str,
    current_user: models.User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    chat_session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        session_id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    messages = (
        db.query(models.ChatMessage)
        .filter(models.ChatMessage.chat_session_id == chat_session.id)
        .order_by(models.ChatMessage.timestamp.asc())
        .all()
    )

    return {
        "session_id": session_id,
        "messages": [
            {
                "user_message": msg.user_message,
                "ai_response": msg.ai_response,
                "timestamp": msg.timestamp.isoformat() + "Z",
            }
            for msg in messages
        ],
    }

@router.post("/save_chat_message")
def save_chat_message(
    message_data: ChatMessageSave,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    chat_session = db.query(models.ChatSession).filter(
        models.ChatSession.id == message_data.chat_id,
        models.ChatSession.user_id == current_user.id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    existing = db.query(models.ChatMessage).filter(
        models.ChatMessage.chat_session_id == message_data.chat_id,
        models.ChatMessage.user_message == message_data.user_message,
        models.ChatMessage.ai_response == message_data.ai_response,
    ).first()

    if existing:
        return {"status": "success", "message": "Message already exists"}

    chat_message = models.ChatMessage(
        chat_session_id=message_data.chat_id,
        user_id=current_user.id,
        user_message=message_data.user_message,
        ai_response=message_data.ai_response,
        is_user=True,
    )
    db.add(chat_message)

    chat_session.updated_at = datetime.now(timezone.utc)

    message_count = db.query(models.ChatMessage).filter(
        models.ChatMessage.chat_session_id == message_data.chat_id
    ).count()

    if chat_session.title == "New Chat" and message_count <= 1:
        _apply_fallback_chat_title(chat_session, message_data.user_message)

    try:
        from services.gamification_system import award_points
        award_points(db, chat_session.user_id, "ai_chat")
    except Exception:
        pass

    try:
        from activity_logger import log_activity
        log_activity(
            user_id=chat_session.user_id,
            tool_name="ai_chat",
            action="send_message",
            tokens_used=0,
            metadata={"chat_id": message_data.chat_id},
        )
    except Exception:
        pass

    db.commit()
    return {"status": "success", "message": "Message saved successfully"}

@router.delete("/delete_chat_session/{session_id}")
def delete_chat_session(
    session_id: str,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    chat_session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        session_id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    session_pk = chat_session.id
    db.query(models.MessageMLLog).filter(
        models.MessageMLLog.session_id == session_pk
    ).delete(synchronize_session=False)
    db.query(models.CerbylSessionState).filter(
        models.CerbylSessionState.session_id == session_pk
    ).delete(synchronize_session=False)
    db.query(models.ChatConceptSignal).filter(
        models.ChatConceptSignal.chat_session_id == session_pk
    ).delete(synchronize_session=False)
    db.query(models.ChatTutorState).filter(
        models.ChatTutorState.chat_session_id == session_pk
    ).delete(synchronize_session=False)
    db.query(models.ChatMessage).filter(
        models.ChatMessage.chat_session_id == session_pk
    ).delete(synchronize_session=False)
    db.delete(chat_session)
    db.commit()

    return {"status": "success"}

@router.post("/submit_response_feedback")
async def submit_response_feedback(
    user_id: str = Form(...),
    rating: int = Form(...),
    message_context: str = Form(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    user = current_user

    feedback = models.UserFeedback(
        user_id=user.id,
        feedback_type="rating",
        rating=rating,
        topic_context=message_context,
        is_processed=False,
    )
    db.add(feedback)
    db.commit()

    return {"status": "success", "message": "Feedback recorded"}

@router.post("/submit_advanced_feedback")
async def submit_advanced_feedback(
    user_id: str = Form(...),
    rating: int = Form(...),
    feedback_text: str = Form(None),
    improvement_suggestion: str = Form(None),
    message_content: str = Form(None),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    user = current_user

    feedback = models.UserFeedback(
        user_id=user.id,
        feedback_type="rating_with_feedback",
        rating=rating,
        feedback_text=feedback_text,
        topic_context=message_content,
        is_processed=False,
    )
    db.add(feedback)
    db.commit()

    return {"status": "success", "message": "Feedback recorded"}

@router.post("/generate_chat_title")
async def generate_chat_title(
    request: GenerateChatTitleRequest,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(request.user_id, current_user)
    chat_session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        request.chat_id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    messages = (
        db.query(models.ChatMessage)
        .filter(models.ChatMessage.chat_session_id == chat_session.id)
        .order_by(models.ChatMessage.timestamp.asc())
        .limit(5)
        .all()
    )

    if not messages:
        return {"title": "New Chat", "status": "no_messages"}

    fallback_title = _fallback_chat_title(messages[0].user_message)

    conversation = "\n".join(
        [f"Student: {m.user_message}\nTutor: {m.ai_response}" for m in messages[:3]]
    )

    prompt = (
        f"Generate a short, descriptive title (3-6 words) for this conversation:\n\n"
        f"{conversation}\n\nTitle:"
    )

    try:
        title = (await call_ai_async(prompt, max_tokens=30, temperature=0.7)).strip().strip('"').strip("'")
        title = re.sub(r"\s+", " ", title).strip(" .,:;!?")
        if not title or title.lower() in {"new chat", "chat session", "untitled chat"}:
            title = fallback_title
        if len(title) > 50:
            title = title[:47] + "..."

        chat_session.title = title
        db.commit()

        return {"title": title, "status": "success"}
    except Exception as e:
        logger.error(f"Error generating title: {e}")
        if fallback_title != "New Chat":
            chat_session.title = fallback_title
            db.commit()
        return {"title": fallback_title, "status": "fallback"}

@router.post("/generate_chat_summary")
async def generate_chat_summary(
    chat_id: str = Form(...),
    user_id: str = Form(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    chat_session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        chat_id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    messages = (
        db.query(models.ChatMessage)
        .filter(models.ChatMessage.chat_session_id == chat_session.id)
        .order_by(models.ChatMessage.timestamp.asc())
        .all()
    )

    if not messages:
        return {"summary": "No messages in this chat.", "status": "empty"}

    conversation = "\n".join(
        [f"Student: {m.user_message}\nTutor: {m.ai_response}" for m in messages[-10:]]
    )

    prompt = (
        f"Summarize this tutoring conversation in 2-3 sentences:\n\n{conversation}\n\nSummary:"
    )

    try:
        summary = await call_ai_async(prompt, max_tokens=200, temperature=0.5)
        return {"summary": summary.strip(), "status": "success"}
    except Exception as e:
        return {"summary": f"Could not generate summary: {e}", "status": "error"}

@router.get("/check_proactive_message")
async def check_proactive_message(
    user_id: str = Query(...),
    db: Session = Depends(get_db),
):
    return {"has_message": False, "message": None}

@router.post("/generate_welcome_message")
async def generate_welcome_message(
    user_id: str = Form(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    user = current_user
    name = user.first_name if user else "there"
    return {
        "message": f"Hey {name}! What would you like to learn today?",
        "status": "success",
    }

@router.get("/conversation_starters")
async def get_conversation_starters(
    user_id: str = Query(None),
    db: Session = Depends(get_db),
):
    return {
        "starters": [
            "Explain a concept I'm struggling with",
            "Help me prepare for an exam",
            "Quiz me on a topic",
            "Summarize my recent notes",
        ]
    }

@router.post("/create_chat_folder")
def create_chat_folder(
    folder_data: ChatFolderCreate,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(folder_data.user_id, current_user)
    user = current_user

    folder = models.ChatFolder(
        user_id=user.id,
        name=folder_data.name,
        color=folder_data.color,
        parent_id=folder_data.parent_id,
    )
    db.add(folder)
    db.commit()
    db.refresh(folder)

    return {
        "id": folder.id,
        "name": folder.name,
        "color": folder.color,
        "status": "success",
    }

@router.get("/get_chat_folders")
def get_chat_folders(
    user_id: str = Query(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    user = current_user

    folders = db.query(models.ChatFolder).filter(models.ChatFolder.user_id == user.id).all()

    return {
        "folders": [
            {
                "id": f.id,
                "name": f.name,
                "color": f.color,
                "parent_id": f.parent_id,
                "created_at": f.created_at.isoformat() + "Z" if f.created_at else None,
            }
            for f in folders
        ]
    }

@router.delete("/delete_chat_folder/{folder_id}")
def delete_chat_folder(
    folder_id: int,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    folder = db.query(models.ChatFolder).filter(
        models.ChatFolder.id == folder_id,
        models.ChatFolder.user_id == current_user.id,
    ).first()
    if not folder:
        raise HTTPException(status_code=404, detail="Folder not found")

    db.query(models.ChatSession).filter(
        models.ChatSession.folder_id == folder_id
    ).update({"folder_id": None})
    db.delete(folder)
    db.commit()

    return {"status": "success"}

@router.put("/move_chat_to_folder")
def move_chat_to_folder(
    data: ChatUpdateFolder,
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    chat = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        data.chat_id,
    ).first()
    if not chat:
        raise HTTPException(status_code=404, detail="Chat session not found")

    if data.folder_id is not None:
        folder = db.query(models.ChatFolder).filter(
            models.ChatFolder.id == data.folder_id,
            models.ChatFolder.user_id == current_user.id,
        ).first()
        if not folder:
            raise HTTPException(status_code=404, detail="Folder not found")

    chat.folder_id = data.folder_id
    db.commit()

    return {"status": "success"}

@router.post("/convert_chat_to_note_content/")
async def convert_chat_to_note_content(
    chat_id: str = Form(...),
    user_id: str = Form(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    chat_session = resolve_by_id_or_uid(
        db.query(models.ChatSession).filter(models.ChatSession.user_id == current_user.id),
        models.ChatSession,
        chat_id,
    ).first()
    if not chat_session:
        raise HTTPException(status_code=404, detail="Chat session not found")

    messages = (
        db.query(models.ChatMessage)
        .filter(models.ChatMessage.chat_session_id == chat_session.id)
        .order_by(models.ChatMessage.timestamp.asc())
        .all()
    )

    if not messages:
        raise HTTPException(status_code=404, detail="No messages found")

    conversation = "\n\n".join(
        [f"**Q:** {m.user_message}\n\n**A:** {m.ai_response}" for m in messages]
    )

    prompt = (
        f"Convert this Q&A conversation into well-organized study notes with headers and key points:\n\n"
        f"{conversation}\n\nStudy Notes:"
    )

    try:
        notes_content = await call_ai_async(prompt, max_tokens=2000, temperature=0.5)
        return {"content": notes_content.strip(), "status": "success"}
    except Exception as e:
        return {"content": conversation, "status": "fallback"}

@router.post("/ai_group_notes")
async def ai_group_notes(
    user_id: str = Form(...),
    db: Session = Depends(get_db),
    current_user: models.User = Depends(get_current_user),
):
    _assert_user_matches_request(user_id, current_user)
    user = current_user

    notes = (
        db.query(models.Note)
        .filter(models.Note.user_id == user.id, models.Note.is_deleted == False)
        .all()
    )

    if not notes:
        return {"groups": [], "status": "no_notes"}

    notes_text = "\n".join([f"- {n.title}" for n in notes[:50]])

    prompt = (
        f"Group these notes into logical categories. Return JSON array of objects "
        f'with "folder_name" and "note_titles" fields:\n\n{notes_text}'
    )

    try:
        result = await call_ai_async(prompt, max_tokens=500, temperature=0.5)
        import json
        groups = json.loads(_strip_json_code_fence(result.strip()))
        return {"groups": groups, "status": "success"}
    except Exception:
        return {"groups": [], "status": "error"}
